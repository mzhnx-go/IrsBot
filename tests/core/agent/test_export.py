"""对话导出渲染层单测（纯函数，不依赖数据库 / HTTP / LLM）。

覆盖：
- extract_text：content 的三种形态（dict / str / 其它）
- build_filename：非法字符清洗、超长截断、空标题兜底、扩展名拼装
- 各格式渲染内容与结构：Markdown / TXT / JSON / Word / PDF
- 只导出 user / assistant —— system / tool 属内部产物，不得出现在导出文件里
- render_export 统一入口：文件名、Content-Type、非法格式抛 ValueError
"""

import io
import json
import uuid
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from app.core.agent.export import (
    ExportFormat,
    build_filename,
    extract_text,
    render_docx,
    render_export,
    render_json,
    render_markdown,
    render_pdf,
    render_txt,
)

#: 固定导出时刻，避免断言里出现「当前时间」导致结果不可复现
EXPORTED_AT = datetime(2026, 9, 20, 3, 0, 0, tzinfo=timezone.utc)


def _conversation(title: str = "财务报表分析") -> SimpleNamespace:
    """造一个 Conversation 替身（只用到渲染需要的 4 个字段）。"""
    return SimpleNamespace(
        id=uuid.UUID("11111111-2222-3333-4444-555555555555"),
        title=title,
        created_at=datetime(2026, 9, 19, 10, 0, 0, tzinfo=timezone.utc),
        updated_at=datetime(2026, 9, 19, 10, 5, 0, tzinfo=timezone.utc),
    )


def _message(role: str, text: str, offset_seconds: int = 0) -> SimpleNamespace:
    """造一个 Message 替身；offset_seconds 用于错开时间戳。"""
    return SimpleNamespace(
        role=role,
        content={"text": text},
        created_at=datetime(2026, 9, 19, 10, 0, 0, tzinfo=timezone.utc)
        + timedelta(seconds=offset_seconds),
    )


def _sample_messages() -> list[SimpleNamespace]:
    """一段含 user / assistant / system / tool 的完整样本。"""
    return [
        _message("system", "你是财务助手", 0),
        _message("user", "帮我看看这个报表", 1),
        _message("assistant", "先看\n现金流", 2),
        _message("tool", "工具原始返回", 3),
    ]


# ── extract_text ──────────────────────────────────────────────


def test_extract_text_from_dict() -> None:
    """当前落库形态 {"text": ...} 取 text 键"""
    assert extract_text({"text": "你好"}) == "你好"


def test_extract_text_from_str() -> None:
    """兼容早期直接存字符串的数据"""
    assert extract_text("直接存的字符串") == "直接存的字符串"


@pytest.mark.parametrize("raw", [None, [], {}])
def test_extract_text_empty_returns_blank(raw) -> None:
    """空值（None / 空列表 / 空字典）→ 空串，绝不抛异常"""
    assert extract_text(raw) == ""


def test_extract_text_unknown_type_json_fallback() -> None:
    """既不是 str 也不是 dict 的非空值 → JSON 序列化兜底（脏数据不能把导出打挂）"""
    assert extract_text(123) == "123"
    assert extract_text([1, 2]) == "[1, 2]"


def test_extract_text_dict_without_text_key() -> None:
    """dict 里没有 text 键 → 空串（不做「随便取第一个值」的猜测）"""
    assert extract_text({"image": "xxx"}) == ""


# ── build_filename ────────────────────────────────────────────


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("财务报表分析", "财务报表分析.pdf"),
        ("a/b\\c:d*e?f", "a_b_c_d_e_f.pdf"),
        ("  空白标题  ", "空白标题.pdf"),
        ("", "对话.pdf"),
        ("   ", "对话.pdf"),
        ("...", "对话.pdf"),
    ],
)
def test_build_filename_sanitizes(title: str, expected: str) -> None:
    """标题里的路径分隔符/通配符/控制字符一律替换，空标题回落「对话」"""
    assert build_filename(title, ExportFormat.PDF) == expected


def test_build_filename_truncates_long_title() -> None:
    """超长标题截断（主体 60 字 + "." + 扩展名），避免撑爆文件系统路径上限"""
    name = build_filename("字" * 500, ExportFormat.MD)
    assert name.endswith(".md")
    assert name == "字" * 60 + ".md"


def test_build_filename_extension_per_format() -> None:
    """每种格式拼自己的扩展名"""
    for fmt, ext in ExportFormat.EXTENSION.items():
        assert build_filename("t", fmt).endswith(f".{ext}")


# ── Markdown ──────────────────────────────────────────────────


def test_markdown_structure_and_roles() -> None:
    """Markdown：一级标题 + 元信息 + 每条消息的角色小节与时间"""
    out = render_markdown(
        _conversation(), _sample_messages(), exported_at=EXPORTED_AT
    )

    assert out.startswith("# 财务报表分析\n")
    assert "- 导出时间：2026-09-20 03:00:00 UTC" in out
    assert "- 消息条数：2" in out
    assert "## 用户 · 2026-09-19 10:00:01 UTC" in out
    assert "## 助手 · 2026-09-19 10:00:02 UTC" in out
    assert "帮我看看这个报表" in out
    assert "先看\n现金流" in out


def test_markdown_excludes_system_and_tool() -> None:
    """system 系统提示词与 tool 原始返回不得出现在导出内容里"""
    out = render_markdown(
        _conversation(), _sample_messages(), exported_at=EXPORTED_AT
    )
    assert "你是财务助手" not in out
    assert "工具原始返回" not in out


def test_markdown_empty_conversation() -> None:
    """空会话也能导出，给出明确占位而不是空白文件"""
    out = render_markdown(_conversation(), [], exported_at=EXPORTED_AT)
    assert "消息条数：0" in out
    assert "（该对话暂无消息）" in out


def test_markdown_keeps_markup_verbatim() -> None:
    """Markdown 不做转义：用户内容里的 # 保留原样（它就是标记语言本身）"""
    msgs = [_message("user", "# 这是用户写的标题")]
    out = render_markdown(_conversation(), msgs, exported_at=EXPORTED_AT)
    assert "# 这是用户写的标题" in out


# ── TXT ───────────────────────────────────────────────────────


def test_txt_has_no_markdown_marks() -> None:
    """纯文本里不出现 Markdown 标记符号"""
    out = render_txt(_conversation(), _sample_messages(), exported_at=EXPORTED_AT)

    assert out.startswith("财务报表分析\n")
    assert "#" not in out
    assert "[用户] 2026-09-19 10:00:01 UTC" in out
    assert "[助手] 2026-09-19 10:00:02 UTC" in out


def test_txt_empty_conversation() -> None:
    """空会话占位"""
    assert "（该对话暂无消息）" in render_txt(
        _conversation(), [], exported_at=EXPORTED_AT
    )


# ── JSON ──────────────────────────────────────────────────────


def test_json_is_parseable_and_structured() -> None:
    """JSON 可被解析，字段语义完整、消息顺序与时间正序一致"""
    raw = render_json(_conversation(), _sample_messages(), exported_at=EXPORTED_AT)
    data = json.loads(raw)

    assert data["conversation"]["title"] == "财务报表分析"
    assert data["conversation"]["id"] == "11111111-2222-3333-4444-555555555555"
    assert data["conversation"]["created_at"] == "2026-09-19T10:00:00Z"
    assert data["exported_at"] == "2026-09-20T03:00:00Z"
    assert data["message_count"] == 2
    assert [m["role"] for m in data["messages"]] == ["user", "assistant"]
    assert data["messages"][1]["content"] == "先看\n现金流"
    assert data["messages"][0]["created_at"] == "2026-09-19T10:00:01Z"


def test_json_keeps_chinese_unescaped() -> None:
    """中文原样输出，不写成 \\uXXXX 转义（人也要能直接读）"""
    raw = render_json(_conversation(), _sample_messages(), exported_at=EXPORTED_AT)
    assert "帮我看看这个报表" in raw
    assert "\\u" not in raw


# ── Word ──────────────────────────────────────────────────────


def test_docx_is_valid_and_contains_turns() -> None:
    """Word 文档可被解析，标题/角色行/正文都在"""
    from docx import Document

    content = render_docx(
        _conversation(), _sample_messages(), exported_at=EXPORTED_AT
    )
    doc = Document(io.BytesIO(content))
    texts = [p.text for p in doc.paragraphs]

    assert any("财务报表分析" == t for t in texts)
    assert any(t.startswith("用户 · ") for t in texts)
    assert any(t.startswith("助手 · ") for t in texts)
    assert "帮我看看这个报表" in texts


def test_docx_preserves_multiline_content() -> None:
    """多行回复必须还原成换行，不能被 Word 吞成一行"""
    from docx import Document

    msgs = [_message("assistant", "第一行\n第二行\n第三行")]
    doc = Document(io.BytesIO(render_docx(_conversation(), msgs, exported_at=EXPORTED_AT)))
    assert "第一行\n第二行\n第三行" in [p.text for p in doc.paragraphs]


def test_docx_empty_conversation() -> None:
    """空会话占位"""
    from docx import Document

    doc = Document(io.BytesIO(render_docx(_conversation(), [], exported_at=EXPORTED_AT)))
    assert any("（该对话暂无消息）" == p.text for p in doc.paragraphs)


# ── PDF ───────────────────────────────────────────────────────


def test_pdf_is_valid_and_text_extractable() -> None:
    """PDF 结构合法，且中文能被原样抽回（证明字体映射正确，不是方框）"""
    from pypdf import PdfReader

    content = render_pdf(_conversation(), _sample_messages(), exported_at=EXPORTED_AT)

    assert content.startswith(b"%PDF")
    text = PdfReader(io.BytesIO(content)).pages[0].extract_text()
    assert "财务报表分析" in text
    assert "帮我看看这个报表" in text


def test_pdf_escapes_markup_characters() -> None:
    """正文里的 & < > 必须转义 —— 否则 reportlab 按迷你 HTML 解析会直接抛错"""
    msgs = [_message("user", "帮我写个 <div> & <span> 标签")]
    content = render_pdf(_conversation(), msgs, exported_at=EXPORTED_AT)
    assert content.startswith(b"%PDF")

    from pypdf import PdfReader

    text = PdfReader(io.BytesIO(content)).pages[0].extract_text()
    assert "<div>" in text and "&" in text


def test_pdf_uses_bundled_font_asset() -> None:
    """PDF 必须用仓库随包分发的中文字体。

    这条是回归防线：字体路径算错时不会报错，只会**静默**退回
    不嵌入字体的 STSong-Light —— 本机能看、缺中文字体的环境一片方框。
    """
    from app.core.agent import export as export_module

    assert export_module._PDF_FONT_PATH.is_file(), (
        f"字体文件缺失或路径算错：{export_module._PDF_FONT_PATH}"
    )
    assert export_module._pdf_font_name() == export_module._PDF_FONT_NAME


def test_pdf_embeds_font_subset() -> None:
    """PDF 里必须真的嵌入字体子集（FontFile2），否则渲染全看阅读器脸色"""
    content = render_pdf(_conversation(), _sample_messages(), exported_at=EXPORTED_AT)
    assert b"FontFile2" in content


def test_pdf_empty_conversation() -> None:
    """空会话也能生成合法 PDF"""
    assert render_pdf(_conversation(), [], exported_at=EXPORTED_AT).startswith(b"%PDF")


# ── render_export 统一入口 ────────────────────────────────────


@pytest.mark.parametrize("fmt", ExportFormat.ALL)
def test_render_export_returns_file_metadata(fmt: str) -> None:
    """每种格式都返回「文件名 + 非空字节 + 对应 Content-Type」"""
    exported = render_export(
        _conversation(), _sample_messages(), fmt, exported_at=EXPORTED_AT
    )

    assert exported.filename == f"财务报表分析.{ExportFormat.EXTENSION[fmt]}"
    assert exported.content
    assert exported.media_type == ExportFormat.MEDIA_TYPE[fmt]


def test_render_export_rejects_unknown_format() -> None:
    """未知格式抛 ValueError（路由层已用 Literal 挡在前面，这里是二次防线）"""
    with pytest.raises(ValueError, match="不支持的导出格式"):
        render_export(_conversation(), [], "exe", exported_at=EXPORTED_AT)


def test_render_export_defaults_exported_at_to_now() -> None:
    """不传 exported_at 时取当前 UTC 时间（保证调用方不会漏传就报错）"""
    before = datetime.now(timezone.utc) - timedelta(seconds=5)
    exported = render_export(_conversation(), [], ExportFormat.MD)
    assert before <= datetime.now(timezone.utc)

    text = exported.content.decode("utf-8")
    year = str(datetime.now(timezone.utc).year)
    assert f"- 导出时间：{year}-" in text